"""Importación de un XLSX de banco de punta a punta (SPEC-013, corrección 2026-09-29).

Cubre lo que el servicio hace con el formato nuevo: persistir el extracto y sus
movimientos, que la cuenta 572 la elija el usuario (el XLSX trae IBAN, no codigo
del plan), que la deduplicacion por `sha256` siga valiendo, que un extracto
descuadrado no deje nada a medias, y que un extracto de otra empresa no se vea.

El constructor del XLSX esta en `tests/unit/extracto_xlsx_support.py`, con numeros
inventados: el extracto real de `Data/` no se copia al repositorio.
"""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from models.treasury.extracto_bancario import ExtractoBancario
from models.treasury.movimiento_bancario import MovimientoBancario
from services.reconciliation.importacion import ImportacionError, importar_extracto
from tests.conftest import sembrar_empresa_pgc
from tests.unit.extracto_xlsx_support import xlsx_banco


async def test_importa_el_xlsx_y_persista_sus_movimientos(
    db_session: AsyncSession,
) -> None:
    await sembrar_empresa_pgc(db_session, 10)
    extracto = await importar_extracto(
        db_session,
        empresa_id=10,
        file_bytes=xlsx_banco(),
        nombre_fichero="movimientos ene_feb26.xlsx",
        layout="xlsx_bancario",
        cuenta_codigo="5720",
        actor="t",
    )
    assert extracto.n_movimientos == 4
    assert extracto.saldo_inicial == Decimal("1863.7400")
    assert extracto.saldo_final == Decimal("4623.7400")
    n = await db_session.scalar(
        select(func.count(MovimientoBancario.id)).where(
            MovimientoBancario.extracto_id == extracto.id
        )
    )
    assert n == 4


async def test_los_movimientos_guardan_el_signo_y_el_importe_en_positivo(
    db_session: AsyncSession,
) -> None:
    """`importe > 0` es un CHECK de la tabla: el signo va en la columna `signo`."""
    await sembrar_empresa_pgc(db_session, 10)
    extracto = await importar_extracto(
        db_session, empresa_id=10, file_bytes=xlsx_banco(),
        nombre_fichero="x.xlsx", layout="xlsx_bancario", cuenta_codigo="5720",
    )
    filas = (
        (
            await db_session.execute(
                select(MovimientoBancario)
                .where(MovimientoBancario.extracto_id == extracto.id)
                .order_by(MovimientoBancario.orden)
            )
        )
        .scalars()
        .all()
    )
    assert [(f.signo.value, f.importe) for f in filas] == [
        ("H", Decimal("1125.0000")),
        ("H", Decimal("1805.0000")),
        ("H", Decimal("425.0000")),
        ("D", Decimal("595.0000")),
    ]
    assert all(f.importe > 0 for f in filas)


async def test_el_iban_queda_en_la_traza_de_auditoria(db_session: AsyncSession) -> None:
    """No hay columna para el IBAN: la unica traza de que cuenta se subio es el
    `payload` del `audit_log`, asi que ahi tiene que estar."""
    import json

    from models.audit.audit_log import AuditLog

    await sembrar_empresa_pgc(db_session, 10)
    await importar_extracto(
        db_session, empresa_id=10, file_bytes=xlsx_banco(),
        nombre_fichero="x.xlsx", layout="xlsx_bancario", cuenta_codigo="5720",
    )
    filas = (
        (
            await db_session.execute(
                select(AuditLog).where(
                    AuditLog.entidad == "extracto_bancario",
                    AuditLog.operacion == "IMPORT_EXTRACTO",
                )
            )
        )
        .scalars()
        .all()
    )
    assert len(filas) == 1
    payload = json.loads(filas[0].payload or "{}")
    assert payload["iban"] == "ES9121000418450200051332"
    assert payload["layout"] == "xlsx_bancario"
    assert payload["saldo_inicial"] == "1863.7400"


async def test_sin_cuenta_no_importa_y_dice_cual_es_el_iban(
    db_session: AsyncSession,
) -> None:
    """El XLSX no trae codigo del plan, y el error tiene que ayudar."""
    await sembrar_empresa_pgc(db_session, 10)
    with pytest.raises(ImportacionError) as exc:
        await importar_extracto(
            db_session, empresa_id=10, file_bytes=xlsx_banco(),
            nombre_fichero="x.xlsx", layout="xlsx_bancario",
        )
    assert exc.value.code == "cuenta_requerida"
    assert "ES91" in str(exc.value)
    assert await db_session.scalar(select(func.count(ExtractoBancario.id))) == 0


async def test_la_cuenta_indicada_tiene_que_existar_en_la_empresa_activa(
    db_session: AsyncSession,
) -> None:
    await sembrar_empresa_pgc(db_session, 10)
    with pytest.raises(ImportacionError) as exc:
        await importar_extracto(
            db_session, empresa_id=10, file_bytes=xlsx_banco(),
            nombre_fichero="x.xlsx", layout="xlsx_bancario", cuenta_codigo="5799",
        )
    assert exc.value.code == "cuenta_no_encontrada"
    assert await db_session.scalar(select(func.count(MovimientoBancario.id))) == 0


async def test_reimportar_el_mismo_xlsx_da_duplicado(db_session: AsyncSession) -> None:
    await sembrar_empresa_pgc(db_session, 10)
    datos = xlsx_banco()
    await importar_extracto(
        db_session, empresa_id=10, file_bytes=datos, nombre_fichero="a.xlsx",
        layout="xlsx_bancario", cuenta_codigo="5720",
    )
    with pytest.raises(ImportacionError) as exc:
        await importar_extracto(
            db_session, empresa_id=10, file_bytes=datos, nombre_fichero="b.xlsx",
            layout="xlsx_bancario", cuenta_codigo="5720",
        )
    assert exc.value.code == "extracto_duplicado"
    assert await db_session.scalar(select(func.count(ExtractoBancario.id))) == 1


async def test_un_extracto_descuadrado_no_persiste_nada(db_session: AsyncSession) -> None:
    await sembrar_empresa_pgc(db_session, 10)
    libro = openpyxl.Workbook()
    hoja = libro.active
    for columna, nombre in enumerate(["Fecha Operación", "Importe", "Saldo"], start=1):
        hoja.cell(row=1, column=columna, value=nombre)
    hoja.cell(row=2, column=1, value="02/01/2026")
    hoja.cell(row=2, column=2, value=100.0)
    hoja.cell(row=2, column=3, value=500.0)
    hoja.cell(row=3, column=1, value="03/01/2026")
    hoja.cell(row=3, column=2, value=-100.0)
    hoja.cell(row=3, column=3, value=900.0)  # no cuadra: 500 + (-100) = 400
    buffer = io.BytesIO()
    libro.save(buffer)
    with pytest.raises(ImportacionError) as exc:
        await importar_extracto(
            db_session, empresa_id=10, file_bytes=buffer.getvalue(),
            nombre_fichero="roto.xlsx", layout="xlsx_bancario", cuenta_codigo="5720",
        )
    assert exc.value.code == "layout_invalido"
    assert await db_session.scalar(select(func.count(ExtractoBancario.id))) == 0
    assert await db_session.scalar(select(func.count(MovimientoBancario.id))) == 0


async def test_un_extracto_en_dolares_no_se_importa(db_session: AsyncSession) -> None:
    await sembrar_empresa_pgc(db_session, 10)
    with pytest.raises(ImportacionError) as exc:
        await importar_extracto(
            db_session, empresa_id=10, file_bytes=xlsx_banco(divisa="USD"),
            nombre_fichero="usd.xlsx", layout="xlsx_bancario", cuenta_codigo="5720",
        )
    assert exc.value.code == "divisa_no_soportada"
    assert await db_session.scalar(select(func.count(ExtractoBancario.id))) == 0


async def test_un_xlsx_de_otra_empresa_no_se_ve(db_session: AsyncSession) -> None:
    """Constitución III: lo que importa la empresa 10 no existe para la 20."""
    await sembrar_empresa_pgc(db_session, 10)
    await sembrar_empresa_pgc(db_session, 20)
    await importar_extracto(
        db_session, empresa_id=10, file_bytes=xlsx_banco(),
        nombre_fichero="x.xlsx", layout="xlsx_bancario", cuenta_codigo="5720",
    )
    assert await db_session.scalar(
        select(func.count(ExtractoBancario.id)).where(ExtractoBancario.empresa_id == 10)
    ) == 1
    assert await db_session.scalar(
        select(func.count(ExtractoBancario.id)).where(ExtractoBancario.empresa_id == 20)
    ) == 0
    assert await db_session.scalar(
        select(func.count(MovimientoBancario.id)).where(MovimientoBancario.empresa_id == 20)
    ) == 0


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------


def test_el_endpoint_importa_el_xlsx(recon_client) -> None:
    client, token, _ = recon_client
    resp = client.post(
        "/api/v1/extractos",
        files={
            "file": (
                "movimientos.xlsx",
                xlsx_banco(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"layout": "xlsx_bancario", "cuenta": "5720"},
        headers={"Authorization": f"Bearer {token}", "X-Empresa-Activa": "10"},
    )
    assert resp.status_code == 201, resp.text
    cuerpo = resp.json()
    assert cuerpo["n_movimientos"] == 4
    assert cuerpo["saldo_inicial"] == "1863.7400"
    assert cuerpo["saldo_final"] == "4623.7400"
    assert cuerpo["fecha_inicio"] == "2026-01-02"
    assert cuerpo["fecha_fin"] == "2026-02-27"


def test_un_formato_que_no_existe_responde_422_con_los_que_si(
    recon_client,
) -> None:
    """Con el despacho silencioso de antes, un `layout` mal escrito caia en el
    parser de ancho fijo y contestaba 'longitud 22 != 100': un error de formato
    hablando del ancho de un fichero que no es de ancho fijo."""
    client, token, _ = recon_client
    resp = client.post(
        "/api/v1/extractos",
        files={"file": ("x.xlsx", xlsx_banco(), "application/octet-stream")},
        data={"layout": "xlsx", "cuenta": "5720"},
        headers={"Authorization": f"Bearer {token}", "X-Empresa-Activa": "10"},
    )
    assert resp.status_code == 422, resp.text
    detalle = resp.json()["detail"]
    assert detalle["error"] == "layout_desconocido"
    assert "xlsx_bancario" in detalle["soportados"]


def test_sin_cuenta_responde_422_diciendo_el_iban(recon_client) -> None:
    client, token, _ = recon_client
    resp = client.post(
        "/api/v1/extractos",
        files={"file": ("x.xlsx", xlsx_banco(), "application/octet-stream")},
        data={"layout": "xlsx_bancario"},
        headers={"Authorization": f"Bearer {token}", "X-Empresa-Activa": "10"},
    )
    assert resp.status_code == 422, resp.text
    assert resp.json()["detail"]["error"] == "cuenta_requerida"
    assert "ES91" in resp.json()["detail"]["detail"]


def test_un_xlsx_descuadrado_responde_422(recon_client) -> None:
    client, token, _ = recon_client
    resp = client.post(
        "/api/v1/extractos",
        files={
            "file": (
                "x.xlsx",
                xlsx_banco([("02/01/2026", "02/01/2026", "Uno", 100.0, 500.0, "072", None),
                            ("03/01/2026", "03/01/2026", "Dos", -100.0, 900.0, "072", None)]),
                "application/octet-stream",
            )
        },
        data={"layout": "xlsx_bancario", "cuenta": "5720"},
        headers={"Authorization": f"Bearer {token}", "X-Empresa-Activa": "10"},
    )
    assert resp.status_code == 422, resp.text
    assert resp.json()["detail"]["error"] == "layout_invalido"
    assert client.get(
        "/api/v1/extractos", headers={"Authorization": f"Bearer {token}", "X-Empresa-Activa": "10"}
    ).json()["total"] == 0
